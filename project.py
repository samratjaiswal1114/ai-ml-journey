import openpyxl as xl

wb = xl.load_workbook('AXIS.xlsx')
sheet = wb['Sheet1']
cell = sheet['a1']
cell = sheet.cell(1,1)

for row in range(2,10):
    cell = sheet.cell(row,2)
    print(cell.value)